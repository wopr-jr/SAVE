from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from enum import Enum
from pathlib import Path
from uuid import UUID
from typing import Any, TYPE_CHECKING

if TYPE_CHECKING:
    from SAVE.modules.exporter.interface import ExportOptions


EXCEL_MAX_CELL_LENGTH = 32_767


SEVERITY_FILLS = {
    "critical": "7030A0",
    "high": "FF0000",
    "medium": "FFC000",
    "low": "92D050",
    "unknown": "BFBFBF",
}


STATUS_FILLS = {
    "open": "FF9999",
    "not_a_finding": "C6EFCE",
    "not_applicable": "D9EAD3",
    "not_reviewed": "FFF2CC",
}


FINDINGS_HEADERS = (
    "Checklist UUID",
    "STIG ID",
    "STIG Name",
    "Vuln ID",
    "Group ID",
    "Group Title",
    "Rule ID",
    "Rule Version",
    "Severity",
    "Status",
    "Rule Title",
    "Comments",
    "Finding Details",
    "Discussion",
    "Check Content",
    "Fix Text",
    "CCI References",
    "Legacy IDs",
    "Target Key",
    "Created At",
    "Updated At",
)


STIG_HEADERS = (
    "STIG UUID",
    "STIG ID",
    "STIG Name",
    "Display Name",
    "Release Info",
    "Reference Identifier",
    "Rule Count",
)


ASSET_HEADERS = (
    "Property",
    "Value",
)


def export_xlsx(
    checklist: Any,
    destination: str | Path,
    *,
    options: ExportOptions | None = None,
) -> list[str]:
    """
    Export a normalized SAVE Checklist as an XLSX workbook.

    Workbook sheets:
        Summary
        Asset
        STIGs
        Findings

    Returns:
        A list of non-fatal export warnings.
    """
    options = options or ExportOptions()
    destination_path = Path(destination)
    warnings: list[str] = []

    Workbook, Font, PatternFill, Alignment, Border, Side, Table, TableStyleInfo = (
        _load_openpyxl()
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    _write_summary_sheet(
        worksheet=summary_sheet,
        checklist=checklist,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
        Border=Border,
        Side=Side,
    )

    _write_asset_sheet(
        workbook=workbook,
        checklist=checklist,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
        Border=Border,
        Side=Side,
    )

    _write_stigs_sheet(
        workbook=workbook,
        checklist=checklist,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
        Table=Table,
        TableStyleInfo=TableStyleInfo,
        warnings=warnings,
    )

    _write_findings_sheet(
        workbook=workbook,
        checklist=checklist,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
        Table=Table,
        TableStyleInfo=TableStyleInfo,
        warnings=warnings,
    )

    workbook.save(destination_path)

    return warnings


def _xlsx_export_handler(
    checklist: Any,
    destination: Path,
    options: ExportOptions,
) -> list[str]:
    """
    Handler signature expected by SAVE's unified ExportService registry.
    """
    return export_xlsx(
        checklist,
        destination,
        options=options,
    )


def _write_summary_sheet(
    *,
    worksheet,
    checklist: Any,
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
) -> None:
    """
    Write a readable high-level checklist summary.
    """
    title_font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )

    heading_font = Font(
        bold=True,
        color="FFFFFF",
    )

    title_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    heading_fill = PatternFill(
        "solid",
        fgColor="4472C4",
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="BFBFBF",
        )
    )

    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "SAVE Checklist Summary"
    worksheet["A1"].font = title_font
    worksheet["A1"].fill = title_fill
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    worksheet["A3"] = "Checklist Metadata"
    worksheet["A3"].font = heading_font
    worksheet["A3"].fill = heading_fill

    metadata = (
        ("Title", _get_attr(checklist, "title")),
        (
            "Checklist UUID",
            _get_attr(checklist, "checklist_uuid"),
        ),
        (
            "Checklist Format",
            _enum_value(
                _get_attr(
                    checklist,
                    "checklist_format",
                )
            ),
        ),
        (
            "Checklist Version",
            _get_attr(
                checklist,
                "checklist_version",
            ),
        ),
        (
            "Source Filename",
            _get_attr(
                checklist,
                "source_filename",
            ),
        ),
        (
            "Source SHA-256",
            _get_attr(
                checklist,
                "source_sha256",
            ),
        ),
        (
            "Imported At",
            _get_attr(
                checklist,
                "imported_at",
            ),
        ),
    )

    row = 4

    for label, value in metadata:
        worksheet.cell(
            row=row,
            column=1,
            value=label,
        ).font = Font(bold=True)

        worksheet.cell(
            row=row,
            column=2,
            value=_excel_value(
                value,
                warnings=[],
                context=label,
            ),
        )

        worksheet.cell(
            row=row,
            column=1,
        ).border = thin_border

        worksheet.cell(
            row=row,
            column=2,
        ).border = thin_border

        row += 1

    status_counts = _status_counts(checklist)
    severity_counts = _severity_counts(checklist)

    worksheet["A13"] = "Finding Status Summary"
    worksheet["A13"].font = heading_font
    worksheet["A13"].fill = heading_fill

    status_rows = (
        ("Open", status_counts["open"]),
        (
            "Not a Finding",
            status_counts["not_a_finding"],
        ),
        (
            "Not Applicable",
            status_counts["not_applicable"],
        ),
        (
            "Not Reviewed",
            status_counts["not_reviewed"],
        ),
    )

    row = 14

    for label, count in status_rows:
        worksheet.cell(
            row=row,
            column=1,
            value=label,
        )

        worksheet.cell(
            row=row,
            column=2,
            value=count,
        )

        row += 1

    worksheet["D13"] = "Severity Summary"
    worksheet["D13"].font = heading_font
    worksheet["D13"].fill = heading_fill

    severity_rows = (
        ("Critical", severity_counts["critical"]),
        ("High", severity_counts["high"]),
        ("Medium", severity_counts["medium"]),
        ("Low", severity_counts["low"]),
        ("Unknown", severity_counts["unknown"]),
    )

    row = 14

    for label, count in severity_rows:
        worksheet.cell(
            row=row,
            column=4,
            value=label,
        )

        worksheet.cell(
            row=row,
            column=5,
            value=count,
        )

        row += 1

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 55
    worksheet.column_dimensions["C"].width = 5
    worksheet.column_dimensions["D"].width = 24
    worksheet.column_dimensions["E"].width = 15


def _write_asset_sheet(
    *,
    workbook,
    checklist: Any,
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
) -> None:
    """
    Write checklist target-asset metadata.
    """
    worksheet = workbook.create_sheet(
        title="Asset",
    )

    _write_headers(
        worksheet=worksheet,
        headers=ASSET_HEADERS,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
    )

    asset = _get_attr(
        checklist,
        "asset",
    )

    asset_rows = (
        (
            "Target Type",
            _get_attr(
                asset,
                "target_type",
            ),
        ),
        (
            "Host Name",
            _get_attr(
                asset,
                "host_name",
            ),
        ),
        (
            "FQDN",
            _get_attr(
                asset,
                "fqdn",
            ),
        ),
        (
            "IP Address",
            _get_attr(
                asset,
                "ip_address",
            ),
        ),
        (
            "MAC Address",
            _get_attr(
                asset,
                "mac_address",
            ),
        ),
        (
            "Target Key",
            _get_attr(
                asset,
                "target_key",
            ),
        ),
        (
            "Role",
            _get_attr(
                asset,
                "role",
            ),
        ),
        (
            "Technology Area",
            _get_attr(
                asset,
                "technology_area",
            ),
        ),
        (
            "Comments",
            _get_attr(
                asset,
                "comments",
            ),
        ),
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9D9D9",
        )
    )

    for row_index, (label, value) in enumerate(
        asset_rows,
        start=2,
    ):
        worksheet.cell(
            row=row_index,
            column=1,
            value=label,
        ).font = Font(bold=True)

        worksheet.cell(
            row=row_index,
            column=2,
            value=_excel_value(
                value,
                warnings=[],
                context=f"Asset {label}",
            ),
        )

        worksheet.cell(
            row=row_index,
            column=1,
        ).border = thin_border

        worksheet.cell(
            row=row_index,
            column=2,
        ).border = thin_border

    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 75


def _write_stigs_sheet(
    *,
    workbook,
    checklist: Any,
    Font,
    PatternFill,
    Alignment,
    Table,
    TableStyleInfo,
    warnings: list[str],
) -> None:
    """
    Write one STIG row per normalized Stig object.
    """
    worksheet = workbook.create_sheet(
        title="STIGs",
    )

    _write_headers(
        worksheet=worksheet,
        headers=STIG_HEADERS,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
    )

    stigs = _get_attr(
        checklist,
        "stigs",
        default=[],
    ) or []

    for stig in stigs:
        worksheet.append(
            (
                _excel_value(
                    _get_attr(
                        stig,
                        "stig_uuid",
                        "uuid",
                    ),
                    warnings=warnings,
                    context="STIG UUID",
                ),
                _excel_value(
                    _get_attr(
                        stig,
                        "stig_id",
                    ),
                    warnings=warnings,
                    context="STIG ID",
                ),
                _excel_value(
                    _get_attr(
                        stig,
                        "stig_name",
                    ),
                    warnings=warnings,
                    context="STIG Name",
                ),
                _excel_value(
                    _get_attr(
                        stig,
                        "display_name",
                    ),
                    warnings=warnings,
                    context="STIG Display Name",
                ),
                _excel_value(
                    _get_attr(
                        stig,
                        "release_info",
                    ),
                    warnings=warnings,
                    context="STIG Release Info",
                ),
                _excel_value(
                    _get_attr(
                        stig,
                        "reference_identifier",
                    ),
                    warnings=warnings,
                    context="STIG Reference Identifier",
                ),
                len(
                    _get_attr(
                        stig,
                        "rules",
                        default=[],
                    ) or []
                ),
            )
        )

    _finish_table_sheet(
        worksheet=worksheet,
        table_name="StigsTable",
        Table=Table,
        TableStyleInfo=TableStyleInfo,
        wrap_columns={
            2,
            3,
            4,
            5,
            6,
        },
    )


def _write_findings_sheet(
    *,
    workbook,
    checklist: Any,
    Font,
    PatternFill,
    Alignment,
    Table,
    TableStyleInfo,
    warnings: list[str],
) -> None:
    """
    Write one spreadsheet row per normalized rule/finding.
    """
    worksheet = workbook.create_sheet(
        title="Findings",
    )

    _write_headers(
        worksheet=worksheet,
        headers=FINDINGS_HEADERS,
        Font=Font,
        PatternFill=PatternFill,
        Alignment=Alignment,
    )

    checklist_uuid = _get_attr(
        checklist,
        "checklist_uuid",
    )

    stigs = _get_attr(
        checklist,
        "stigs",
        default=[],
    ) or []

    for stig in stigs:
        rules = _get_attr(
            stig,
            "rules",
            default=[],
        ) or []

        for rule in rules:
            severity = _enum_value(
                _get_attr(
                    rule,
                    "severity",
                )
            )

            status = _enum_value(
                _get_attr(
                    rule,
                    "status",
                )
            )

            worksheet.append(
                (
                    _excel_value(
                        checklist_uuid,
                        warnings=warnings,
                        context="Checklist UUID",
                    ),
                    _excel_value(
                        _get_attr(
                            stig,
                            "stig_id",
                        ),
                        warnings=warnings,
                        context="STIG ID",
                    ),
                    _excel_value(
                        _get_attr(
                            stig,
                            "stig_name",
                            "display_name",
                        ),
                        warnings=warnings,
                        context="STIG Name",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "vuln_id",
                            "vuln_num",
                        ),
                        warnings=warnings,
                        context="Vuln ID",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "group_id",
                            "group_id_source",
                            "group_id_src",
                        ),
                        warnings=warnings,
                        context="Group ID",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "group_title",
                        ) or _leaf_group_title(rule),
                        warnings=warnings,
                        context="Group Title",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "rule_id",
                            "rule_id_source",
                            "rule_id_src",
                        ),
                        warnings=warnings,
                        context="Rule ID",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "rule_version",
                        ),
                        warnings=warnings,
                        context="Rule Version",
                    ),
                    _excel_value(
                        severity,
                        warnings=warnings,
                        context="Severity",
                    ),
                    _excel_value(
                        status,
                        warnings=warnings,
                        context="Status",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "title",
                            "rule_title",
                        ),
                        warnings=warnings,
                        context="Rule Title",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "comments",
                        ),
                        warnings=warnings,
                        context="Comments",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "finding_details",
                        ),
                        warnings=warnings,
                        context="Finding Details",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "discussion",
                        ),
                        warnings=warnings,
                        context="Discussion",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "check_content",
                            "checkcontent",
                        ),
                        warnings=warnings,
                        context="Check Content",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "fix_text",
                            "fixtext",
                        ),
                        warnings=warnings,
                        context="Fix Text",
                    ),
                    _excel_value(
                        _join_values(
                            _get_attr(
                                rule,
                                "ccis",
                                "cci_refs",
                                default=[],
                            ) or []
                        ),
                        warnings=warnings,
                        context="CCI References",
                    ),
                    _excel_value(
                        _join_values(
                            _get_attr(
                                rule,
                                "legacy_ids",
                                default=[],
                            ) or []
                        ),
                        warnings=warnings,
                        context="Legacy IDs",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "target_key",
                        ),
                        warnings=warnings,
                        context="Target Key",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "created_at",
                        ),
                        warnings=warnings,
                        context="Created At",
                    ),
                    _excel_value(
                        _get_attr(
                            rule,
                            "updated_at",
                        ),
                        warnings=warnings,
                        context="Updated At",
                    ),
                )
            )

    _finish_table_sheet(
        worksheet=worksheet,
        table_name="FindingsTable",
        Table=Table,
        TableStyleInfo=TableStyleInfo,
        wrap_columns={
            3,
            6,
            11,
            12,
            13,
            14,
            15,
            16,
            17,
            18,
        },
    )

    _apply_findings_colors(
        worksheet=worksheet,
        PatternFill=PatternFill,
    )


def _write_headers(
    *,
    worksheet,
    headers: tuple[str, ...],
    Font,
    PatternFill,
    Alignment,
) -> None:
    """
    Write a consistently styled worksheet header row.
    """
    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=header,
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def _finish_table_sheet(
    *,
    worksheet,
    table_name: str,
    Table,
    TableStyleInfo,
    wrap_columns: set[int],
) -> None:
    """
    Apply table, filter, freeze-pane, wrap, and width behavior.
    """
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = cell.alignment.copy(
                vertical="top",
                wrap_text=cell.column in wrap_columns,
            )

    _auto_size_columns(worksheet)

    if worksheet.max_row < 2:
        return

    from openpyxl.utils import get_column_letter

    last_column = get_column_letter(
        worksheet.max_column
    )

    table = Table(
        displayName=table_name,
        ref=f"A1:{last_column}{worksheet.max_row}",
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    worksheet.add_table(table)


def _apply_findings_colors(
    *,
    worksheet,
    PatternFill,
) -> None:
    """
    Apply visual highlighting to Severity and Status columns.

    Findings column layout:
        I = Severity
        J = Status
    """
    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        severity_cell = worksheet.cell(
            row=row_index,
            column=9,
        )

        status_cell = worksheet.cell(
            row=row_index,
            column=10,
        )

        severity_value = str(
            severity_cell.value or ""
        ).strip().lower()

        status_value = str(
            status_cell.value or ""
        ).strip().lower()

        severity_color = SEVERITY_FILLS.get(
            severity_value,
        )

        if severity_color:
            severity_cell.fill = PatternFill(
                "solid",
                fgColor=severity_color,
            )

        status_color = STATUS_FILLS.get(
            status_value,
        )

        if status_color:
            status_cell.fill = PatternFill(
                "solid",
                fgColor=status_color,
            )


def _auto_size_columns(
    worksheet,
) -> None:
    """
    Set practical worksheet column widths.

    Long text fields are capped at 60 characters and use wrapped text.
    """
    from openpyxl.utils import get_column_letter

    for column_cells in worksheet.columns:
        max_length = 0
        column_index = column_cells[0].column

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            max_length = max(
                max_length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(
            max(max_length + 2, 12),
            60,
        )


def _status_counts(
    checklist: Any,
) -> Counter[str]:
    counts: Counter[str] = Counter()

    for stig in _get_attr(
        checklist,
        "stigs",
        default=[],
    ) or []:
        for rule in _get_attr(
            stig,
            "rules",
            default=[],
        ) or []:
            status = _enum_value(
                _get_attr(
                    rule,
                    "status",
                )
            ) or "not_reviewed"

            counts[status.lower()] += 1

    return counts


def _severity_counts(
    checklist: Any,
) -> Counter[str]:
    counts: Counter[str] = Counter()

    for stig in _get_attr(
        checklist,
        "stigs",
        default=[],
    ) or []:
        for rule in _get_attr(
            stig,
            "rules",
            default=[],
        ) or []:
            severity = _enum_value(
                _get_attr(
                    rule,
                    "severity",
                )
            ) or "unknown"

            counts[severity.lower()] += 1

    return counts


def _get_attr(
    source: Any,
    *names: str,
    default: Any = None,
) -> Any:
    """
    Return the first non-None matching model attribute.

    Supports field aliases while the SAVE model evolves.
    """
    if source is None:
        return default

    for name in names:
        value = getattr(
            source,
            name,
            None,
        )

        if value is not None:
            return value

    return default


def _leaf_group_title(
    rule: Any,
) -> str | None:
    group_tree = _get_attr(
        rule,
        "group_tree",
        default=[],
    ) or []

    if not group_tree:
        return None

    return _get_attr(
        group_tree[-1],
        "title",
    )


def _join_values(
    values: list[Any],
) -> str | None:
    if not values:
        return None

    return "\n".join(
        str(value)
        for value in values
        if value is not None
    )


def _enum_value(
    value: Any,
) -> str | None:
    if value is None:
        return None

    if isinstance(value, Enum):
        return str(value.value)

    return str(value)


def _excel_value(
    value: Any,
    *,
    warnings: list[str],
    context: str,
) -> Any:
    """
    Convert SAVE model values into Excel-supported values.

    Excel cells cannot contain more than 32,767 characters. Long fields such
    as Check Content, Fix Text, and Finding Details are truncated safely.
    """
    if value is None:
        return None

    if isinstance(value, Enum):
        value = value.value

    elif isinstance(value, UUID):
        value = str(value)

    elif isinstance(value, (datetime, date)):
        value = value.isoformat()

    elif isinstance(value, bool):
        return value

    elif isinstance(value, (int, float)):
        return value

    else:
        value = str(value)

    if not isinstance(value, str):
        return value

    if len(value) <= EXCEL_MAX_CELL_LENGTH:
        return value

    warnings.append(
        f"{context} exceeded Excel's {EXCEL_MAX_CELL_LENGTH:,}-character "
        "cell limit and was truncated."
    )

    return (
        value[: EXCEL_MAX_CELL_LENGTH - 80]
        + f"\n\n[TRUNCATED BY SAVE; original length={len(value)}]"
    )


def _load_openpyxl():
    """
    Delay the optional dependency import until XLSX export is requested.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Alignment,
            Border,
            Font,
            PatternFill,
            Side,
        )
        from openpyxl.worksheet.table import (
            Table,
            TableStyleInfo,
        )

    except ImportError as exc:
        raise RuntimeError(
            "XLSX export requires the optional dependency openpyxl. "
            'Install it with: pip install "openpyxl>=3.1,<4"'
        ) from exc

    return (
        Workbook,
        Font,
        PatternFill,
        Alignment,
        Border,
        Side,
        Table,
        TableStyleInfo,
    )